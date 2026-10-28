# -*- coding: utf-8 -*-
"""生成演示素材：3 页 PPT(两版) + 2 页 Word(两版) + 位图(两版) + 代码(两版)。

只用来给 README 出效果图，内容为虚构演示数据，不含任何真实文档内容。
页数刻意压到 3 / 2：对比图才不会变成一堆小图。
"""
import os
import sys

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Pt as DPt, RGBColor as DRGB
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = sys.argv[1]
os.makedirs(OUT, exist_ok=True)

NAVY = RGBColor(0x1B, 0x2A, 0x3A)
RED = RGBColor(0x9E, 0x2A, 0x25)
CREAM = RGBColor(0xFD, 0xF6, 0xE6)
GREY = RGBColor(0x3A, 0x3A, 0x3A)
MUTE = RGBColor(0x7C, 0x86, 0x92)
LINE = RGBColor(0xE2, 0xD9, 0xC4)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xC9, 0xD4, 0xE0)

W, H = Inches(13.333), Inches(7.5)
ML = 0.95          # 左边距
CW = 13.333 - 0.95 * 2   # 内容宽 11.433


def blank(prs, bg):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(1, 0, 0, W, H)
    r.fill.solid()
    r.fill.fore_color.rgb = bg
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def text(slide, x, y, w, h, s, size, color, bold=False, align=PP_ALIGN.LEFT, spacing=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = 'Microsoft YaHei'
    if spacing:
        p.line_spacing = spacing
    return tb


def rect(slide, x, y, w, h, color, shape=1, line=None, lw=1):
    r = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if color is None:
        r.fill.background()
    else:
        r.fill.solid()
        r.fill.fore_color.rgb = color
    if line is None:
        r.line.fill.background()
    else:
        r.line.color.rgb = line
        r.line.width = Pt(lw)
    r.shadow.inherit = False
    return r


def header(slide, eyebrow, title, on_dark=False):
    """统一页眉：小标签 + 大标题 + 细线。"""
    text(slide, ML, 0.72, 8.0, 0.4, eyebrow, 13, RED, bold=True)
    text(slide, ML, 1.12, 11.0, 0.85, title, 34, WHITE if on_dark else NAVY, bold=True)
    rect(slide, ML, 2.08, 1.5, 0.045, RED)


def cover_graphic(path):
    """封面右侧的抽象图形（演示图片渲染能力）。"""
    w, h = 818, 1000
    im = Image.new('RGB', (w, h), (27, 42, 58))
    dr = ImageDraw.Draw(im)
    for i in range(9):
        x = 60 + i * 78
        dr.rectangle([x, 180 + (i % 4) * 46, x + 44, 760], fill=(36 + i * 5, 62 + i * 5, 92 + i * 4))
    dr.ellipse([210, 300, 610, 700], fill=(158, 42, 37))
    dr.ellipse([300, 390, 520, 610], fill=(196, 154, 92))
    dr.rectangle([470, 120, 700, 210], fill=(253, 246, 230))
    im.save(path)
    return path


def build_pptx(path, version):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    v2 = (version == 'v2')
    title = '版本对比交付' + ('（修订版）' if v2 else '')

    # ---------- 1 封面 ----------
    s = blank(prs, NAVY)
    gpath = cover_graphic(os.path.join(OUT, '_cover_graphic.png'))
    s.shapes.add_picture(gpath, Inches(7.55), 0, height=H)
    rect(s, 7.5, 0, 0.055, 7.5, CREAM)

    text(s, ML, 1.25, 6.2, 0.4, 'MINIMAX CODE  ·  MAVIS  SKILL', 12.5, RED, bold=True)
    text(s, ML, 1.8, 6.4, 1.95, title, 52, WHITE, bold=True)
    rect(s, ML, 3.85, 1.4, 0.05, RED)
    text(s, ML, 4.18, 6.2, 0.6, '改文件 · 留原版 · 出对比图', 21, PALE)
    for i, chip in enumerate(['原版不动', '递增版本', '可视对比']):
        x = ML + i * 1.62
        rect(s, x, 5.05, 1.42, 0.44, None, line=RGBColor(0x46, 0x5A, 0x70), lw=1)
        text(s, x, 5.16, 1.42, 0.3, chip, 12.5, PALE, align=PP_ALIGN.CENTER)
    text(s, ML, 6.4, 6.2, 0.4, 'github.com/H-u-iCN/version-diff-delivery', 12, RGBColor(0x76, 0x86, 0x99))

    # ---------- 2 四个能力 ----------
    s = blank(prs, CREAM)
    header(s, '能力  /  CAPABILITIES', '把「改了什么」变成看得见的产物')
    cards = [
        ('01', '原版不动', 'V0 永远留在原处，任何时候都能回到起点'),
        ('02', '递增版本', '改动落在 _V1 / _V2，重名自动跳号'),
        ('03', '对比图', '代码 / 图片 / PPT / Word 各有对应的可视形式'),
        ('04', '一行安装', 'git clone 进 skills 目录，下个会话即生效'),
    ]
    cw, ch, gap = 5.47, 1.72, 0.5
    for i, (num, t, b) in enumerate(cards):
        r, c = divmod(i, 2)
        x = ML + c * (cw + gap)
        y = 2.55 + r * (ch + 0.42)
        rect(s, x, y, cw, ch, WHITE, line=LINE, lw=1)
        rect(s, x, y, 0.06, ch, RED)
        rect(s, x + 0.45, y + 0.42, 0.46, 0.46, NAVY, shape=MSO_SHAPE.OVAL)
        text(s, x + 0.45, y + 0.53, 0.46, 0.3, num, 14, CREAM, bold=True, align=PP_ALIGN.CENTER)
        text(s, x + 1.18, y + 0.38, cw - 1.6, 0.5, t, 21, NAVY, bold=True)
        text(s, x + 1.18, y + 0.92, cw - 1.6, 0.6, b, 13, MUTE)

    # ---------- 3 三步流程 ----------
    s = blank(prs, CREAM)
    header(s, '流程  /  WORKFLOW', '一轮迭代，三个动作')
    steps = [
        ('01', '查基线', '脚本判定上一版是谁\n下一版叫什么，不手算'),
        ('02', '开新版本', '以基线为蓝本拷贝一份\n原版与旧版本原地不动'),
        ('03', '渲染对比', '按类型选渲染模式\n输出到 .mavis-compare/'),
    ]
    bw = 3.35
    for i, (num, t, b) in enumerate(steps):
        x = ML + i * (bw + 0.59)
        rect(s, x, 2.7, bw, 2.45, WHITE, line=LINE, lw=1)
        text(s, x + 0.42, 2.95, 1.6, 0.75, num, 40, RED, bold=True)
        text(s, x + 0.42, 3.85, bw - 0.8, 0.45, t, 20, NAVY, bold=True)
        text(s, x + 0.42, 4.35, bw - 0.8, 0.75, b, 12.5, MUTE, spacing=1.35)
        if i < 2:
            rect(s, x + bw + 0.12, 3.72, 0.35, 0.42, RED, shape=MSO_SHAPE.CHEVRON)
    rect(s, ML, 5.55, CW, 0.9, NAVY)
    text(s, ML + 0.5, 5.83, CW - 1.0, 0.4,
         '收尾固定给：版本链表格  +  每个改动文件一张对比图  +  改动摘要',
         15, CREAM, align=PP_ALIGN.CENTER)
    prs.save(path)
    return path


def build_docx(path, version):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Microsoft YaHei'
    st.font.size = DPt(10.5)

    def para(s, size=10.5, bold=False, color=None, align=None, space_after=4, space_before=0):
        p = doc.add_paragraph()
        r = p.add_run(s)
        r.font.size = DPt(size)
        r.font.bold = bold
        if color:
            r.font.color.rgb = DRGB(*color)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == 'c' else (
            WD_ALIGN_PARAGRAPH.RIGHT if align == 'r' else WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.space_after = DPt(space_after)
        p.paragraph_format.space_before = DPt(space_before)
        return p

    def h1(s):
        para(s, 15, True, (0x9E, 0x2A, 0x25), space_after=6, space_before=10)

    para('版本对比交付 Skill' + ('（修订版）' if version == 'v2' else ''), 22, True,
         (0x1B, 0x2A, 0x3A), 'c', 6)
    para('—— 演示文档：每一版如何被留痕、被对比', 11, False, (0x7C, 0x86, 0x92), 'c', 16)

    h1('一、它解决什么问题')
    para('日常改文件最怕两件事：一是把原版改没了，二是说"改了什么"全靠嘴。这里把两件事都变成可复核的产物。', 10.5, space_after=8)
    for t in ['原版是 V0，字节不动，任何时候都能回到起点',
              '改动落在新版本文件上，命名规则固定，不靠记忆',
              '对比图与汇总表同时给出：图看位置，表看清单']:
        para('· ' + t, 10.5, space_after=3)

    h1('二、版本链长什么样')
    tb = doc.add_table(rows=4, cols=4)
    tb.style = 'Table Grid'
    for ci, v in enumerate(['版本', '文件', '改动', '对比形态']):
        c = tb.cell(0, ci)
        c.text = v
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = DPt(9.5)
    for ri, row in enumerate([('V0', 'config.ts', '原始状态', '—'),
                              ('V1', 'config_V1.ts', '标题文案', '并排 diff 图'),
                              ('V2', 'config_V2.ts', '新增常量', '并排 diff 图')], start=1):
        for ci, v in enumerate(row):
            c = tb.cell(ri, ci)
            c.text = v
            c.paragraphs[0].runs[0].font.size = DPt(9.5)

    doc.add_page_break()
    h1('三、每种文件类型怎么比')
    tb2 = doc.add_table(rows=5, cols=3)
    tb2.style = 'Table Grid'
    for ci, v in enumerate(['文件类型', '主形式', '附带形式']):
        c = tb2.cell(0, ci)
        c.text = v
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.size = DPt(9.5)
    for ri, row in enumerate([('代码 / 配置', '左右并排 diff 图', 'unified diff 代码块'),
                              ('图片', '并排图 + 变更区域框', '尺寸与色值参数表'),
                              ('PPT / Word', '整份接触印相表', '新旧逐页并排 + 橙框'),
                              ('表格', '单元格 before/after', 'CSV 文本 diff')], start=1):
        for ci, v in enumerate(row):
            c = tb2.cell(ri, ci)
            c.text = v
            c.paragraphs[0].runs[0].font.size = DPt(9.5)

    h1('四、装完就能用')
    para('一行命令装进 skills 目录，下个会话生效，不依赖运行时重启：', 10.5, space_after=6)
    para('git clone https://github.com/H-u-iCN/version-diff-delivery.git '
         '~/.minimax/skills/version-diff-delivery', 9.5, False, (0x3A, 0x3A, 0x3A), space_after=10)
    para('依赖只有 Python 3 与 Pillow；PPTX 路线另需 python-pptx，DOCX 路线另需 python-docx。'
         '缺哪个，对应路线会明确报错并给兜底形式。', 10.5, space_after=8)

    h1('五、怎么复现仓库里的效果图')
    para('仓库 docs/ 目录放的全部效果图，都是由本 skill 自己的脚本渲染合成样例生成的：', 10.5, space_after=6)
    for t in ['render_compare.py text / image —— 代码与位图对比',
              'render_deck.py sheet / compare —— PPT 接触印相表与逐页对比',
              'render_doc.py sheet / compare —— Word 接触印相表与逐页对比']:
        para('· ' + t, 10.5, space_after=3)
    para('汇报人：千晖　　2026 年 10 月', 10.5, False, None, 'r', 0, 14)
    doc.save(path)
    return path


def build_cover(path, color):
    im = Image.new('RGB', (1000, 640), (250, 246, 238))
    dr = ImageDraw.Draw(im)
    dr.rectangle([90, 90, 470, 380], fill=color)
    dr.rectangle([560, 120, 910, 250], fill=(60, 62, 70))
    dr.rectangle([560, 300, 910, 470], fill=(225, 218, 200))
    dr.ellipse([120, 430, 440, 570], fill=color)
    im.save(path)
    return path


if __name__ == '__main__':
    build_pptx(os.path.join(OUT, 'demo_v1.pptx'), 'v1')
    build_pptx(os.path.join(OUT, 'demo_v2.pptx'), 'v2')
    build_docx(os.path.join(OUT, 'demo_v1.docx'), 'v1')
    build_docx(os.path.join(OUT, 'demo_v2.docx'), 'v2')
    build_cover(os.path.join(OUT, 'cover_v1.png'), (196, 154, 92))
    build_cover(os.path.join(OUT, 'cover_v2.png'), (37, 99, 235))
    with open(os.path.join(OUT, 'demo_v1.ts'), 'w', encoding='utf-8') as f:
        f.write("// 站点配置\nTITLE = '春季上新'\nSUBTITLE = '副标题'\n\nNAV = ['a', 'b']\n")
    with open(os.path.join(OUT, 'demo_v2.ts'), 'w', encoding='utf-8') as f:
        f.write("// 站点配置\nTITLE = '春季上新·第二版'\nSUBTITLE = '副标题'\nSITE_NAME = '春季上新站'\n\nNAV = ['a', 'b', 'c']\n")
    print('demo assets ready in', OUT)
