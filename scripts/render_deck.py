#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""PPTX / 演示文稿的「接触印相表」渲染与新旧对比。

不依赖 LibreOffice / PowerPoint：用 python-pptx 读结构，Pillow 画页面，
再把整份 deck 拼成一张大图（每页缩略图 + p01/p02 标签），手机点开即可缩放看。

子命令：
  sheet   --pptx <file> --out <sheet.png> [--cols 4] [--thumb-w 440]
          渲染一份 deck，输出全部页的接触印相表

  compare --left <old.pptx> --right <new.pptx> --out <cmp.png> [--thumb-w 440]
          新旧两版逐页并排，变化页打橙框 + 序号，输出成一张大图

  slides  --pptx <file> --outdir <dir>
          逐页导出 PNG（供 render_compare.py pair 模式用）

渲染器思路来自用户工程 经典原著阅读分享/tools/render_preview.py（python-pptx + Pillow
近似渲染），此处通用化：补分组形状、圆角矩形/三角、占位图、渐变兜底、每页溢出检测。
"""

import argparse
import json
import os
import sys
from io import BytesIO

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:  # pragma: no cover
    print(json.dumps({"ok": False, "error": "缺少 Pillow，请先 pip install Pillow"}, ensure_ascii=False))
    sys.exit(2)

try:
    from pptx import Presentation
except ImportError:  # pragma: no cover
    Presentation = None  # 只在跑 PPTX 路线时才需要，本模块可被无 pptx 环境安全导入

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.dont_write_bytecode = True  # 别在 skill 包里留下 __pycache__
from sheetkit import build_compare, build_sheet, font

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # pragma: no cover
    pass

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
P = '{http://schemas.openxmlformats.org/presentationml/2006/main}'

PPI = 110  # 渲染分辨率：像素/英寸。改这里会同时影响整页与缩略图。


def emu_in(v):
    return (v or 0) / 914400.0


def hexrgb(rgb):
    return '#%02X%02X%02X' % (rgb[0], rgb[1], rgb[2])


def rgb_of(color_obj, default=None):
    try:
        if color_obj is None:
            return default
        return hexrgb(color_obj.rgb)
    except Exception:
        return default


def fill_info(shape):
    try:
        spPr = shape._element.find(P + 'spPr')
        if spPr is None:
            return None
        solid = spPr.find(A + 'solidFill')
        if solid is None:
            grad = spPr.find(A + 'gradFill')
            if grad is not None:
                srgb = grad.find('.//' + A + 'srgbClr')
                if srgb is not None:
                    return '#' + srgb.get('val'), 1.0
            return None
        srgb = solid.find(A + 'srgbClr')
        if srgb is None:
            return None
        alpha = 1.0
        a = srgb.find(A + 'alpha')
        if a is not None:
            alpha = int(a.get('val')) / 100000.0
        return '#' + srgb.get('val'), alpha
    except Exception:
        return None


def line_info(shape):
    try:
        ln = shape._element.find('.//' + A + 'ln')
        if ln is None:
            return None
        srgb = ln.find('.//' + A + 'srgbClr')
        if srgb is None:
            return None
        w = ln.get('w')
        lw = max(1, int(int(w) / 12700.0 * PPI / 72.0)) if w else 1
        return '#' + srgb.get('val'), lw
    except Exception:
        return None


def prst_of(shape):
    try:
        g = shape._element.find('.//' + A + 'prstGeom')
        if g is not None:
            return g.get('prst') or 'rect'
    except Exception:
        pass
    return 'rect'


def wrap(draw, text, fnt, max_px):
    lines = []
    for hard in text.split('\n'):
        if not hard:
            lines.append('')
            continue
        cur = ''
        for ch in hard:
            trial = cur + ch
            if draw.textlength(trial, font=fnt) > max_px and cur:
                lines.append(cur)
                cur = ch
            else:
                cur = trial
        lines.append(cur)
    return lines


def para_info(para, default_size=18):
    text = ''.join(r.text for r in para.runs)
    if not text:
        text = para.text
    size = None
    bold = False
    color = '#000000'
    name = 'cn'
    for r in para.runs:
        if r.font.size is not None:
            size = r.font.size.pt
        if r.font.bold:
            bold = True
        try:
            if r.font.name and 'Georgia' in r.font.name:
                name = 'en'
        except Exception:
            pass
        c = rgb_of(r.font.color, None)
        if c:
            color = c
        if size:
            break
    if size is None:
        size = default_size
    kind = ('en_bold' if bold else 'en') if name == 'en' else ('cn_bold' if bold else 'cn')
    align = str(para.alignment) if para.alignment is not None else 'LEFT'
    return text, size, bold, color, kind, align


def draw_text_frame(d, tf, x, y, w, h, default_size=18):
    ml = (tf.margin_left or 0) / 914400.0 * PPI
    mr = (tf.margin_right or 0) / 914400.0 * PPI
    mt = (tf.margin_top or 0) / 914400.0 * PPI
    avail = max(10, w - ml - mr)
    cy = y + mt
    overflow = False
    for para in tf.paragraphs:
        text, size, bold, color, kind, align = para_info(para, default_size)
        if not text:
            cy += 6
            continue
        fnt = font(kind, size * PPI / 72.0)
        for ln in wrap(d, text, fnt, avail):
            tw = d.textlength(ln, font=fnt)
            if 'CENTER' in align:
                tx = x + ml + (avail - tw) / 2.0
            elif 'RIGHT' in align:
                tx = x + ml + (avail - tw)
            else:
                tx = x + ml
            d.text((tx, cy), ln, font=fnt, fill=color)
            cy += size * PPI / 72.0 * 1.24
        cy += 3
    if cy > y + h + 4:
        overflow = True
    return overflow


def draw_shape_fill(canvas, d, shape, x, y, w, h, prst):
    fi = fill_info(shape)
    if not fi:
        lc = line_info(shape)
        if lc and prst != 'line':
            d.rectangle([x, y, x + w, y + h], outline=lc[0], width=lc[1])
        return d
    col, alpha = fi
    hexa = col + '%02X' % int(alpha * 255)
    layer = Image.new('RGBA', canvas.size, (0, 0, 0, 0))
    ld = ImageDraw.Draw(layer)
    if prst == 'ellipse':
        ld.ellipse([x, y, x + w, y + h], fill=hexa)
    elif prst == 'chevron':
        ld.polygon([(x, y), (x + w * 0.6, y), (x + w, y + h / 2),
                    (x + w * 0.6, y + h), (x, y + h), (x + w * 0.4, y + h / 2)], fill=hexa)
    elif prst == 'triangle':
        ld.polygon([(x + w / 2, y), (x + w, y + h), (x, y + h)], fill=hexa)
    elif prst == 'roundRect':
        ld.rounded_rectangle([x, y, x + w, y + h], radius=min(w, h) * 0.12, fill=hexa)
    elif prst == 'line':
        lc = line_info(shape)
        ld.line([x, y, x + w, y + h], fill=(lc[0] if lc else col),
                width=(lc[1] if lc else 1))
    else:
        ld.rectangle([x, y, x + w, y + h], fill=hexa)
    canvas.alpha_composite(layer)
    return ImageDraw.Draw(canvas)


def render_slide(slide, idx, W, H):
    """渲染单页为 RGB 图像；返回 (image, overflow_text_list, errors)。"""
    bg = '#FFFFFF'
    try:
        bg = hexrgb(slide.background.fill.fore_color.rgb)
    except Exception:
        pass
    canvas = Image.new('RGBA', (W, H), bg + 'FF')
    d = ImageDraw.Draw(canvas)
    overflow = []
    errors = []

    def walk(shapes):
        nonlocal d
        for shape in shapes:
            # 定位只兜"取不到 left/top"这一种情况；不要把计算包在 try 里，
            # 否则拼写/未定义这类真 bug 会被 except 吞掉，表现为"整页空白"这种假成功。
            try:
                left, top, width, height = shape.left, shape.top, shape.width, shape.height
            except Exception:
                left = top = width = height = None
            x = emu_in(left) * PPI
            y = emu_in(top) * PPI
            w = emu_in(width) * PPI
            h = emu_in(height) * PPI
            prst = prst_of(shape)

            # 分组形状：递归展开
            if shape.shape_type is not None and str(shape.shape_type).startswith('GROUP'):
                try:
                    walk(shape.shapes)
                    continue
                except Exception:
                    pass

            # 图片 / 图片占位符
            if hasattr(shape, 'image'):
                try:
                    im = Image.open(BytesIO(shape.image.blob)).convert('RGBA')
                    tw, th = im.size
                    if w > 0 and h > 0:
                        tr = w / h
                        if tw / th > tr:
                            nw = int(th * tr)
                            im = im.crop(((tw - nw) // 2, 0, (tw - nw) // 2 + nw, th))
                        else:
                            nh = int(tw / tr)
                            im = im.crop((0, (th - nh) // 2, tw, (th - nh) // 2 + nh))
                        im = im.resize((max(1, int(w)), max(1, int(h))), Image.LANCZOS)
                        canvas.alpha_composite(im, (int(x), int(y)))
                    continue
                except Exception as exc:
                    errors.append('slide %d 图片 %s 渲染失败: %s' % (idx, shape.name, exc))
                    d.rectangle([x, y, x + w, y + h], outline='#FF00FF', width=2)
                    continue

            # 表格
            if getattr(shape, 'has_table', False):
                tbl = shape.table
                colw = [emu_in(c.width) * PPI for c in tbl.columns]
                rowh = [emu_in(r.height) * PPI for r in tbl.rows]
                if not any(rowh):
                    rowh = [h / max(1, len(tbl.rows))] * len(tbl.rows)
                yy = y
                for ri, row in enumerate(tbl.rows):
                    xx = x
                    for ci, cell in enumerate(row.cells):
                        cw = colw[ci] if ci < len(colw) else 60
                        ch = rowh[ri] if ri < len(rowh) else 24
                        cfill = None
                        try:
                            if cell.fill.type is not None and cell.fill.type != 5:
                                cfill = rgb_of(cell.fill.fore_color, None)
                        except Exception:
                            cfill = None
                        if cfill:
                            d.rectangle([xx, yy, xx + cw, yy + ch], fill=cfill)
                        d.rectangle([xx, yy, xx + cw, yy + ch], outline='#8FA6BF', width=1)
                        if cell.text_frame.text.strip():
                            draw_text_frame(d, cell.text_frame, xx, yy, cw, ch, 9)
                        xx += cw
                    yy += rowh[ri] if ri < len(rowh) else 24
                continue

            d = draw_shape_fill(canvas, d, shape, x, y, w, h, prst)

            if getattr(shape, 'has_text_frame', False) and shape.text_frame.text.strip():
                if draw_text_frame(d, shape.text_frame, x, y, w, h):
                    overflow.append(shape.text_frame.text[:34].replace('\n', ' '))
                    d.rectangle([x, y, x + w, y + h], outline='#FF00FF', width=2)

    walk(slide.shapes)
    return canvas.convert('RGB'), overflow, errors


def render_deck(pptx_path, outdir, ppi=None):
    if Presentation is None:
        raise RuntimeError('缺少 python-pptx，请先 pip install python-pptx')
    ppi = ppi or PPI
    if not os.path.isdir(outdir):
        os.makedirs(outdir, exist_ok=True)
    prs = Presentation(pptx_path)
    W = int(prs.slide_width / 914400.0 * ppi)
    H = int(prs.slide_height / 914400.0 * ppi)
    pages, overflow, errors = [], [], []
    for idx, slide in enumerate(prs.slides, start=1):
        img, ov, err = render_slide(slide, idx, W, H)
        path = os.path.join(outdir, 'slide-%02d.png' % idx)
        img.save(path)
        pages.append(path)
        for t in ov:
            overflow.append({'slide': idx, 'text': t})
        errors.extend(err)
    return pages, overflow, errors, (W, H)


# ---------------------------------------------------------------- CLI

def main():
    ap = argparse.ArgumentParser(description="PPTX 接触印相表 / 新旧对比")
    sub = ap.add_subparsers(dest="cmd", required=True)

    s = sub.add_parser("sheet", help="整份 deck 拼成一张接触印相表")
    s.add_argument("--pptx", required=True)
    s.add_argument("--out", required=True)
    s.add_argument("--cols", type=int, default=4)
    s.add_argument("--thumb-w", type=int, default=440)
    s.add_argument("--title", default="")
    s.add_argument("--workdir", default="")

    c = sub.add_parser("compare", help="新旧两版逐页并排对比")
    c.add_argument("--left", required=True)
    c.add_argument("--right", required=True)
    c.add_argument("--out", required=True)
    c.add_argument("--thumb-w", type=int, default=440)
    c.add_argument("--title", default="")
    c.add_argument("--workdir", default="")
    c.add_argument("--no-mark", action="store_true")

    p = sub.add_parser("slides", help="逐页导出 PNG")
    p.add_argument("--pptx", required=True)
    p.add_argument("--outdir", required=True)

    args = ap.parse_args()
    tmp = args.workdir or os.path.join(os.path.dirname(os.path.abspath(args.out)) or ".",
                                       "_deck_work")
    try:
        if args.cmd == "slides":
            pages, overflow, errors, size = render_deck(args.pptx, args.outdir)
            res = {"ok": True, "mode": "slides", "files": pages, "slides": len(pages),
                   "overflow": overflow, "errors": errors, "slide_size": size}
        elif args.cmd == "sheet":
            work = os.path.join(tmp, os.path.splitext(os.path.basename(args.pptx))[0])
            pages, overflow, errors, size = render_deck(args.pptx, work)
            out, dims = build_sheet(pages, args.out, args.cols, args.thumb_w, args.title)
            res = {"ok": True, "mode": "sheet", "files": [out], "slides": len(pages),
                   "overflow": overflow, "errors": errors, "image_size": dims}
        else:
            wl = os.path.join(tmp, "old")
            wr = os.path.join(tmp, "new")
            lp, ov_l, err_l, _ = render_deck(args.left, wl)
            rp, ov_r, err_r, _ = render_deck(args.right, wr)
            out, dims, changed = build_compare(lp, rp, args.out, args.thumb_w, args.title,
                                               mark_changes=not args.no_mark)
            res = {"ok": True, "mode": "compare", "files": [out], "slides": len(rp),
                   "changed_slides": changed, "image_size": dims,
                   "overflow": ov_l + ov_r, "errors": err_l + err_r}
    except Exception as exc:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": "%s: %s" % (type(exc).__name__, exc)},
                         ensure_ascii=False, indent=2))
        sys.exit(1)
    print(json.dumps(res, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
